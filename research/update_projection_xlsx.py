"""Patch GrowthProjection.xlsx in place from research/evidence_current.json
(the workbook builder was lost with the original scratchpad; the layout is
stable, so values are written cell-by-cell). Formulas untouched."""
import json, sys
from pathlib import Path
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1]
ev = json.load(open(ROOT / "research" / "evidence_current.json"))
wb = load_workbook(ROOT / "GrowthProjection.xlsx")
st, data, cfg, cmp_ = wb["Statistics"], wb["Data"], wb["Settings"], wb["Comparison"]
ver = ev["matrix_version"]; w0, w1 = ev["window"]
import datetime as dt
wtxt = f"{dt.date.fromisoformat(w0):%b %Y}–{dt.date.fromisoformat(w1):%b %Y}"
rows = {"CONS": 6, "MOD": 7, "AGG": 8, "VAGG": 9, "SPY": 10, "QQQ": 11}
keys = ["CAGR", "Sortino", "Sharpe", "maxDD", "beta", "y2008", "y2020", "y2022", "y2025", "y2026"]
for name, r in rows.items():
    e = ev["evidence"][name]
    for j, k in enumerate(keys):
        v = e[k]
        st.cell(row=r, column=2 + j, value=round(v, 4) if v is not None else None)
regimes = ["G", "R", "S", "D"]
for i, t in enumerate(["CONS", "MOD", "AGG", "VAGG"]):
    for j, q in enumerate(regimes):
        data.cell(row=6 + i, column=2 + j, value=round(ev["regime_stats"][t][q]["mean_m"], 5))
        data.cell(row=13 + i, column=2 + j, value=round(ev["regime_stats"][t][q]["std_m"], 5))
for j, q in enumerate(regimes):
    data.cell(row=20, column=2 + j, value=round(ev["freq"].get(q, 0.0), 3))
    data.cell(row=21, column=2 + j, value=round(ev["duration_months"].get(q, 0.0), 1))
    cfg.cell(row=6, column=2 + j, value=round(ev["freq"].get(q, 0.0), 2))
# comparison benchmark rows SPY/QQQ (data values)
for name, r in [("SPY", 6), ("QQQ", 7)]:
    e = ev["evidence"][name]
    cmp_.cell(row=r, column=2, value=round(e["CAGR"], 4)); cmp_.cell(row=r, column=3, value=round(e["Sortino"], 2))
    cmp_.cell(row=r, column=4, value=round(e["Sharpe"], 2)); cmp_.cell(row=r, column=5, value=round(e["maxDD"], 4))
    cmp_.cell(row=r, column=6, value=round(e["beta"], 2)); cmp_.cell(row=r, column=7, value=round(e["y2008"], 4)); cmp_.cell(row=r, column=8, value=round(e["y2022"], 4))
    cmp_.cell(row=r + 13, column=2, value=round(e["CAGR"], 4))
# labels
import re
a2 = st["A2"].value or ""
a2 = re.sub(r"^Matrix v\d+ \([^)]*\)", "", a2)
st["A2"] = (f"Matrix {ver} (gold instrument by regime — bullion in Reflation, miners in Deflation for AGG/VAGG; AGG Reflation leverage on the energy leg; "
            f"conditional-duration S, washout-conditional rebound slice, R momentum tilt, regime-gated short book behind a global include-shorts switch)" + a2)
data["A1"] = f"Model Data — per-regime return statistics (matrix {ver}, monthly-only execution, {wtxt})"
cmp_["A2"] = f"Identical window ({wtxt}), dividends included. Benchmark stats are data (blue); tier rows link to Statistics (green)."
db = wb["Dashboard"]
for row in db.iter_rows(min_row=1, max_row=4):
    for c in row:
        if isinstance(c.value, str) and "Backtest" in c.value:
            c.value = re.sub(r"Backtest [A-Za-z]{3} \d{4}–[A-Za-z]{3} \d{4}", f"Backtest {wtxt}", c.value)
wb.calculation.fullCalcOnLoad = True
wb.save(ROOT / "GrowthProjection.xlsx")
print(f"GrowthProjection.xlsx updated to matrix {ver}, window {wtxt}")
for name in ["CONS", "MOD", "AGG", "VAGG"]:
    r = rows[name]; print(f"  {name:>4}:", [st.cell(row=r, column=k).value for k in range(2, 7)])
